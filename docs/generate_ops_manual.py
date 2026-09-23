#!/usr/bin/env python3
"""Generate MA-MinSight operations manual as Word (.docx)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn


def set_cn_font(run, name: str = "宋体", size: int | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)


def add_title_page(doc: Document) -> None:
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MA-MinSight\n安全告警分诊与研判系统")
    r.bold = True
    set_cn_font(r, "黑体", 22)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("操作说明与运维手册")
    set_cn_font(r2, "黑体", 18)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(f"版本：0.1.0    文档日期：{date.today().isoformat()}")
    set_cn_font(r3, "宋体", 12)

    doc.add_page_break()


def add_toc(doc: Document) -> None:
    doc.add_heading("目录", level=1)
    items = [
        "1. 系统概述",
        "2. 系统架构与工作原理",
        "3. 数据处理流水线详解",
        "4. 数据库表说明",
        "5. 部署与初始化",
        "6. Web 控制台操作指南",
        "7. 命令行工具与运维脚本",
        "8. 配置项说明",
        "9. 故障排查手册",
        "10. 附录",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    add_title_page(doc)
    add_toc(doc)

    # ===== 1 概述 =====
    doc.add_heading("1. 系统概述", level=1)
    doc.add_paragraph(
        "MA-MinSight（本项目亦称为 AI-SEC）是一套面向企业安全运营中心（SOC）的告警分诊与 AI 深度调查系统。"
        "系统从 Wazuh Indexer 等上游安全数据源增量拉取告警，经规则初筛、事件聚合、大模型 Agent 调查、"
        "人工协查与防御资产沉淀，形成「告警 → 事件 → 结论 → 判例/白名单」的完整闭环。"
    )
    doc.add_heading("1.1 适用场景", level=2)
    for t in [
        "对接 Wazuh / Suricata 等 SIEM，对海量告警进行自动初筛与降噪；",
        "对高风险或不确定告警启动 AI Agent 深度调查，输出结构化结论与处置建议；",
        "通过企业微信等 IM 通道发起人工协查，补充现场信息；",
        "将已结论事件沉淀为判例、白名单候选、实体画像更新与处置记录，持续优化防御能力。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("1.2 核心设计原则", level=2)
    add_table(
        doc,
        ["原则", "说明"],
        [
            ["初筛 ≠ 最终结论", "triage_results 仅决定路由（归档/深度调查/抽检），Agent 结论在 investigation_conclusions"],
            ["仅白名单可自动归档", "entity_profile_memories 画像记忆不能触发自动放行，只有 active 白名单规则可以"],
            ["Agent 证据驱动", "调查必须调用 Skill 查询资产、历史告警、威胁情报，禁止无证据臆断"],
            ["处置建议不自动执行", "disposition_records 仅为建议，需人工确认后才视为正式处置"],
            ["单租户单机部署", "每台 ECS/VM 独立实例，无多租户隔离字段"],
        ],
    )

    # ===== 2 架构 =====
    doc.add_heading("2. 系统架构与工作原理", level=1)
    doc.add_heading("2.1 逻辑架构", level=2)
    doc.add_paragraph(
        "系统采用 Docker Compose 四服务架构：PostgreSQL（pgvector）存储业务数据；Redis 提供分布式锁与缓存；"
        "app 容器运行 FastAPI 后端与 APScheduler 定时任务；nginx 容器提供 Web 控制台静态页面与 API 反向代理。"
    )
    add_table(
        doc,
        ["组件", "技术", "职责"],
        [
            ["app", "FastAPI + APScheduler", "API、流水线调度、Agent 调查、IM 协查"],
            ["postgres", "PostgreSQL 16 + pgvector", "34 张业务表、配置、审计"],
            ["redis", "Redis 7", "入库/初筛/调查等阶段分布式锁、缓存"],
            ["nginx", "Nginx + React", "HTTPS/HTTP、同源反代 /api、Web 控制台"],
            ["Wazuh Indexer", "OpenSearch 兼容", "上游告警数据源（ECS-B）"],
            ["DeepSeek LLM", "OpenAI 兼容 API", "Agent 调查推理"],
        ],
    )

    doc.add_heading("2.2 数据流总览", level=2)
    doc.add_paragraph(
        "Wazuh Indexer → 告警入库(alerts, NEW) → 初筛分诊(triage_results, TRIAGED/ARCHIVED) → "
        "事件聚合(events, pending_review) → Agent 调查(investigations) → "
        "人工协查(可选) → 防御资产沉淀(判例/白名单/处置) → Event 关闭(closed)。"
    )

    doc.add_heading("2.3 调度机制", level=2)
    doc.add_paragraph(
        "所有流水线阶段由 app 进程内 APScheduler 按 system_config 中的间隔定时触发。"
        "各阶段使用 Redis 全局锁或按数据源锁，防止并发重复处理。"
        "修改调度间隔后需重启 app 容器生效。"
    )
    add_table(
        doc,
        ["调度任务", "默认间隔", "配置键"],
        [
            ["ingestion_poll 告警入库", "5 分钟", "ingestion.poll_interval_minutes"],
            ["triage_poll 初筛分诊", "2 分钟", "triage.poll_interval_minutes"],
            ["aggregation_poll 事件聚合", "2 分钟", "event.aggregation_poll_interval_minutes"],
            ["investigation_poll Agent 调查", "3 分钟", "investigation.poll_interval_minutes"],
            ["human_review_poll 协查派发", "2 分钟", "human_review.poll_interval_minutes"],
            ["defense_assets_poll 防御资产", "5 分钟", "defense_assets.poll_interval_minutes"],
        ],
    )

    # ===== 3 流水线 =====
    doc.add_heading("3. 数据处理流水线详解", level=1)

    stages = [
        (
            "3.1 告警入库（Ingestion）",
            [
                "从 data_sources 表读取已启用数据源（通常为 wazuh），按 cursor_state 游标增量拉取；",
                "Wazuh 适配器使用 search_after 分页，避免深分页性能问题；",
                "告警映射为统一字段写入 alerts 表，source_alert_id 唯一约束去重；",
                "可选 GeoLite2 解析公网 src_ip 归属地，写入 normalized_fields.src_geo；",
                "health_probe 为链路探针专用数据源，不从外部拉取，fetched=0 属正常。",
            ],
        ),
        (
            "3.2 初筛分诊（Triage）",
            [
                "处理 status=NEW 的告警，流水线：去重 → 白名单 → 画像提示 → 规则打分 → 路由 → 持久化；",
                "输出 triage_results，告警变为 TRIAGED（进入深度/不确定队列）或 ARCHIVED（归档）；",
                "archive_whitelist 仅当匹配 status=active 的 whitelist_rules；",
                "profile_hint_score 来自 entity_profiles，不能单独触发归档；",
                "queue_deep_review / queue_uncertain 的告警进入后续聚合。",
            ],
        ),
        (
            "3.3 事件聚合（Aggregation）",
            [
                "将 TRIAGED 且路由为深度/不确定的告警按 host+user+category+24h 窗口合并为 Event；",
                "event_key = SHA256(host|user|category|window_start)；",
                "仅 pending_review 状态的同键 Event 可合并；investigating 等同键会新建 Event；",
                "Event 按 queue_priority、risk_score、last_alert_at 排序供调查消费。",
            ],
        ),
        (
            "3.4 Agent 调查（Investigation）",
            [
                "对 pending_review 的 Event 启动 ReAct 循环，调用 LLM + Skill；",
                "Skills：query_asset（资产/画像）、query_history_alerts（历史告警）、query_threat_intel（威胁情报）、submit_conclusion（提交结论）；",
                "完整轨迹写入 investigation_messages、investigation_tool_calls；",
                "结论写入 investigation_conclusions（verdict/confidence/reasoning/recommended_action）；",
                "每日调查 Event 数受 investigation.daily_event_budget 限制（默认 200）。",
            ],
        ),
        (
            "3.5 人工协查（Human Review）",
            [
                "当 Agent 裁决为 insufficient_information 或 needs_human_review 时，向 IM（企业微信）下发协查问题；",
                "操作员在 Web 协查页或 IM 回复后，系统可触发 Agent 复跑；",
                "协查请求 24 小时超时，最多重发 2 次。",
            ],
        ),
        (
            "3.6 防御资产沉淀（Defense Assets）",
            [
                "对已 concluded 的 Event，从最新 Investigation 沉淀：judgment_cases（判例）、whitelist_candidates（白名单候选）、",
                "profile_update_suggestions（画像更新建议）、disposition_records（处置建议）；",
                "白名单候选经 ≥3 次人工确认后升格为 whitelist_rules.status=active；",
                "默认写入判例后 Event 自动变为 closed。",
            ],
        ),
    ]
    for title, bullets in stages:
        doc.add_heading(title, level=2)
        for b in bullets:
            doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("3.7 推荐手动补跑顺序", level=2)
    doc.add_paragraph(
        "在 Web 设置页手动运行时，请按以下顺序执行："
        "告警入库 → 初筛分诊 → 事件聚合 → Agent 调查 → 协查派发 → 防御资产沉淀。"
        "若某步显示「处理 0 条」，通常是上一步尚无可用数据。"
    )

    doc.add_heading("3.8 初筛路由决策说明", level=2)
    add_table(
        doc,
        ["route_decision", "含义", "告警后续"],
        [
            ["archive_low_risk", "规则分低于低风险阈值", "ARCHIVED，不再进入调查"],
            ["archive_whitelist", "命中 active 白名单", "ARCHIVED"],
            ["archive_dedup", "去重窗口内重复", "ARCHIVED"],
            ["sample_audit", "低风险随机抽检", "ARCHIVED，进入 audit_samples"],
            ["queue_deep_review", "高分，进深度调查", "TRIAGED → 聚合 → Agent"],
            ["queue_uncertain", "分数在不确定区间", "TRIAGED → 聚合 → Agent"],
        ],
    )

    # ===== 4 数据表 =====
    doc.add_heading("4. 数据库表说明", level=1)
    doc.add_paragraph(
        "数据库共 34 张表，以下按业务域分组说明各表用途。所有表使用 UUID 主键（system_config 除外）。"
    )

    table_groups = [
        (
            "4.1 告警接入",
            [
                ("data_sources", "注册上游数据源", "name、adapter_type、config、cursor_state（游标）、is_active"),
                ("alerts", "标准化告警", "source_alert_id、fingerprint、severity、occurred_at、src_ip、host_name、normalized_fields、raw_data、status"),
                ("alert_dedup_groups", "去重窗口分组", "fingerprint、occurrence_count、window_start/end"),
            ],
        ),
        (
            "4.2 初筛与白名单",
            [
                ("triage_results", "初筛结果（不可变）", "alert_id、rule_score、route_decision、route_reason、triaged_at"),
                ("whitelist_rules", "正式白名单规则", "match_pattern、status(active 才可归档)、entity_profile_id"),
                ("whitelist_candidates", "Agent 建议白名单", "suggested_pattern、human_confirmations、status"),
            ],
        ),
        (
            "4.3 实体与威胁情报",
            [
                ("entity_profiles", "资产/主机/IP 画像", "entity_type+entity_key、asset_criticality、is_known_scanner"),
                ("entity_profile_memories", "模糊记忆（不可归档）", "memory_type、content、confidence"),
                ("entity_profile_stats", "滚动统计", "window_days、alert_count、alert_by_category"),
                ("on_duty_knowledge", "值班知识", "scope、content（如内网扫描器 IP）"),
                ("threat_intel_entries", "威胁情报缓存", "ioc_type/value、verdict、source、expires_at"),
            ],
        ),
        (
            "4.4 事件与调查",
            [
                ("events", "调查单元（聚合后）", "event_key、title、status、risk_score、queue_priority、alert_count"),
                ("event_alerts", "Event 与 Alert 关联", "event_id、alert_id"),
                ("investigation_sops", "调查指导原则", "alert_category、guidance_text、recommended_skills"),
                ("investigations", "Agent 调查会话", "event_id、status、model_name、token 用量、skill_call_count"),
                ("investigation_messages", "ReAct 消息历史", "role、content、sequence"),
                ("investigation_tool_calls", "Skill 调用轨迹", "skill_name、input/output、latency_ms"),
                ("investigation_conclusions", "结构化结论", "verdict、confidence、reasoning、recommended_action、human_query"),
            ],
        ),
        (
            "4.5 人工协查",
            [
                ("human_review_requests", "协查请求", "question、channel、status、sent_at"),
                ("human_review_responses", "协查回复", "raw_content、parsed_content、triggered_reinvestigation"),
                ("audit_samples", "低风险抽检", "sample_reason、review_verdict、is_correct"),
            ],
        ),
        (
            "4.6 防御资产",
            [
                ("judgment_cases", "参考判例", "feature_summary、investigation_trace、verdict、embedding"),
                ("rule_candidates", "检测规则草稿", "rule_draft、status"),
                ("profile_update_suggestions", "画像更新建议", "suggested_changes、status"),
                ("disposition_records", "处置建议/确认", "suggested_action、confirmed_action、status"),
            ],
        ),
        (
            "4.7 系统与评测",
            [
                ("users / api_keys", "控制台账号与 API 密钥", "username、is_admin、key_hash"),
                ("system_config", "业务配置键值", "key、value(JSON)、description"),
                ("audit_logs", "系统审计", "actor_type、action、resource_type、detail"),
                ("cache_entries", "PostgreSQL 缓存元数据", "cache_key、cache_type、expires_at"),
                ("regression_test_cases/runs", "Agent 回归测试", "expected_verdict、passed/failed"),
                ("health_check_scenarios/runs", "链路探针", "inject_payload、expected_stage、passed"),
            ],
        ),
    ]

    for group_title, tables in table_groups:
        doc.add_heading(group_title, level=2)
        add_table(doc, ["表名", "用途", "关键字段"], tables)

    doc.add_heading("4.8 告警状态与路由枚举", level=2)
    add_table(
        doc,
        ["枚举", "取值", "含义"],
        [
            ["AlertStatus", "NEW → TRIAGED → EVENT_LINKED / ARCHIVED", "告警生命周期"],
            ["RouteDecision", "archive_* / queue_deep_review / queue_uncertain / sample_audit", "初筛路由"],
            ["EventStatus", "pending_review → investigating → concluded / human_pending → closed", "事件状态"],
            ["InvestigationVerdict", "attack_confirmed / likely_false_positive / insufficient_information / needs_human_review", "Agent 裁决"],
        ],
    )

    # ===== 5 部署 =====
    doc.add_heading("5. 部署与初始化", level=1)
    doc.add_heading("5.1 环境要求", level=2)
    add_table(
        doc,
        ["项目", "最低", "推荐"],
        [
            ["CPU/内存", "2C4G", "4C8G"],
            ["磁盘", "40GB SSD", "80GB SSD"],
            ["操作系统", "Ubuntu 22.04 / Alibaba Cloud Linux 3", "同左"],
            ["软件", "Docker 24+、Docker Compose v2", "同左"],
        ],
    )

    doc.add_heading("5.2 标准部署步骤", level=2)
    for i, step in enumerate(
        [
            "克隆代码到服务器：git clone <仓库> /opt/minsight && cd /opt/minsight",
            "复制环境文件：cp deploy/env/production.env.example .env",
            "编辑 .env：设置 POSTGRES_PASSWORD、API_KEY、JWT_SECRET、ADMIN_INITIAL_PASSWORD、WAZUH_INDEXER_*、DEEPSEEK_API_KEY",
            "（可选）放置 GeoLite2-City.mmdb 到 ./data/geoip/，配置 GEOLITE2_HOST_PATH",
            "构建并启动：docker compose up -d --build",
            "验收：./deploy/scripts/smoke-test.sh",
            "浏览器访问 http://<主机>/ ，使用 admin / ADMIN_INITIAL_PASSWORD 登录并立即改密",
            "初始化资产种子：docker compose exec app ma-entity-seed",
            "（可选）威胁情报同步：docker compose exec app ma-ti-sync",
        ],
        start=1,
    ):
        doc.add_paragraph(f"{i}. {step}")

    doc.add_heading("5.3 ECS-A + ECS-B 典型拓扑", level=2)
    doc.add_paragraph(
        "ECS-A 部署 MA-MinSight（Docker Compose，nginx 6443 HTTPS）；"
        "ECS-B 部署 Wazuh Indexer（端口 20000）。"
        "在 data_sources 配置 indexer_url 指向 ECS-B 内网地址，verify_tls 按证书情况设置。"
        "ECS-A 安全组需放行 6443（控制台）及到 ECS-B 20000 的内网访问。"
    )

    doc.add_heading("5.4 ECS-A HTTPS 启动", level=2)
    for step in [
        "在 .env 设置：COMPOSE_FILE=docker-compose.yml:docker-compose.prod.yml:docker-compose.ecs-a.yml",
        "放置 SSL 证书到 deploy/nginx/certs/fullchain.pem 与 privkey.pem",
        "执行：./deploy/scripts/up-ecs-a.sh",
        "验证：curl -k https://127.0.0.1:6443/health/ready",
    ]:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("5.5 升级与迁移", level=2)
    doc.add_paragraph(
        "git pull 后执行 docker compose build app nginx --no-cache && docker compose up -d --force-recreate app nginx。"
        "app 容器 entrypoint 会自动运行 alembic upgrade head。"
        "修改 .env 后必须 force-recreate app，restart 不会重新加载环境变量。"
    )

    doc.add_heading("5.6 Wazuh 数据源配置", level=2)
    doc.add_paragraph(
        "Wazuh 连接信息存储在 data_sources 表（非 .env 直接配置 URL）。"
        "典型 config JSON 字段：indexer_url、index_pattern（wazuh-alerts-*）、"
        "initial_lookback_hours（游标为空时回溯小时数，默认 24，最大 168）、"
        "verify_tls、username_env/password_env（指向 WAZUH_INDEXER_USER/PASSWORD）。"
    )
    doc.add_paragraph(
        "cursor_state 字段记录增量游标：last_occurred_at（上次拉取到的最新告警时间）、"
        "last_sort_values（OpenSearch search_after 值）、total_ingested（游标计数）。"
    )

    # ===== 6 Web =====
    doc.add_heading("6. Web 控制台操作指南", level=1)
    add_table(
        doc,
        ["菜单", "路径", "功能"],
        [
            ["仪表盘", "/", "运营概览、告警状态分布、流水线最近运行时间、待调查 Event 数"],
            ["事件队列", "/events", "按状态浏览 Event，点击进入详情"],
            ["事件详情", "/events/:id", "关联告警、Agent 结论、协查回复、处置确认"],
            ["调查轨迹", "/investigations/:id/trace", "ReAct 对话流与 Skill 调用记录"],
            ["人工协查", "/human-review", "待回复协查 Inbox、Web 端回复"],
            ["防御资产", "/defense-assets", "白名单确认、画像建议批准、处置确认"],
            ["评测", "/eval", "回归测试与链路探针历史（管理员可运行）"],
            ["设置", "/settings", "故障排查、手动运行流水线、系统配置编辑"],
        ],
    )

    doc.add_heading("6.1 管理员专属功能", level=2)
    for t in [
        "设置 → 手动运行流水线：立即补跑 ingestion/triage/aggregation 等 8 个阶段；",
        "设置 → 系统配置：JSON 编辑 triage 阈值、调查预算、协查开关等；",
        "设置 → 告警入库故障排查：一键检测 Wazuh 连接、游标、回溯窗口；",
        "评测页 → 运行回归测试 / 链路探针。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("6.2 日常运营流程", level=2)
    doc.add_paragraph(
        "正常情况下系统自动每 2–5 分钟调度各阶段，操作员主要关注仪表盘与事件队列。"
        "发现 NEW 告警长期不下降时，检查入库诊断；发现 pending_review Event 堆积时，"
        "检查 Agent 预算与 DEEPSEEK_API_KEY；协查中 Event 需在企业微信或 Web 回复。"
    )

    # ===== 7 CLI =====
    doc.add_heading("7. 命令行工具与运维脚本", level=1)
    doc.add_heading("7.1 容器内 ma-* 命令", level=2)
    add_table(
        doc,
        ["命令", "用途", "示例"],
        [
            ["ma-entity-seed", "种子化资产画像与值班知识", "docker compose exec app ma-entity-seed"],
            ["ma-ti-sync", "同步威胁情报", "docker compose exec app ma-ti-sync --ip x.x.x.x --force"],
            ["ma-geo-enrich", "回填告警 IP 归属地", "docker compose exec app ma-geo-enrich --check"],
        ],
    )

    doc.add_heading("7.2 Python 模块 CLI", level=2)
    add_table(
        doc,
        ["命令", "用途"],
        [
            ["python -m app.ingestion.cli pull --source wazuh", "单数据源手动拉取"],
            ["python -m app.eval.cli health-check run --all", "运行链路探针"],
            ["python -m app.eval.cli regression run", "运行回归测试"],
        ],
    )

    doc.add_heading("7.3 宿主机脚本", level=2)
    add_table(
        doc,
        ["脚本", "用途"],
        [
            ["deploy/scripts/smoke-test.sh", "部署后健康验收"],
            ["deploy/scripts/backup-postgres.sh", "PostgreSQL 备份"],
            ["deploy/scripts/restore-postgres.sh", "从备份恢复"],
            ["deploy/scripts/sync-threat-intel.sh", "威胁情报同步包装"],
            ["deploy/scripts/up-ecs-a.sh", "ECS-A HTTPS 一键启动"],
        ],
    )

    doc.add_heading("7.4 常用 Docker 命令", level=2)
    for cmd in [
        "docker compose logs -f app nginx          # 查看日志",
        "docker compose ps                         # 服务状态",
        "docker compose exec app alembic upgrade head   # 手动迁移",
        "docker compose up -d --force-recreate app      # 重载 .env",
    ]:
        doc.add_paragraph(cmd, style="List Bullet")

    # ===== 8 配置 =====
    doc.add_heading("8. 配置项说明", level=1)
    doc.add_heading("8.1 环境变量（.env）", level=2)
    add_table(
        doc,
        ["变量", "说明"],
        [
            ["POSTGRES_PASSWORD", "数据库密码（必填）"],
            ["WAZUH_INDEXER_USER/PASSWORD", "Wazuh Indexer 凭据"],
            ["DEEPSEEK_API_KEY/MODEL/BASE_URL", "Agent 调查 LLM"],
            ["JWT_SECRET / ADMIN_INITIAL_PASSWORD", "Web 登录"],
            ["WECOM_*", "企业微信协查（6 项）"],
            ["GEOLITE2_CITY_PATH / GEOLITE2_HOST_PATH", "GeoIP 数据库路径"],
            ["ABUSEIPDB_API_KEY / GREYNOISE_API_KEY", "威胁情报 API"],
        ],
    )

    doc.add_heading("8.2 system_config（Web 设置页可编辑）", level=2)
    doc.add_paragraph(
        "业务阈值与调度间隔存储在 system_config 表，管理员可在设置页直接修改 JSON。"
        "关键项包括：triage.deep_review_threshold（默认 70）、investigation.daily_event_budget（默认 200）、"
        "event.aggregation_window_hours（默认 24）、human_review.enabled 等。"
        "完整中文说明见设置页「说明」列。"
    )

    # ===== 9 故障排查 =====
    doc.add_heading("9. 故障排查手册", level=1)

    doc.add_heading("9.1 健康检查端点", level=2)
    add_table(
        doc,
        ["端点", "含义"],
        [
            ["GET /health", "进程存活"],
            ["GET /health/ready", "PostgreSQL + Redis 就绪（Compose healthcheck 使用）"],
        ],
    )

    doc.add_heading("9.2 告警入库 fetched=0", level=2)
    doc.add_paragraph("现象：手动或自动入库显示拉取 0 条。排查步骤：")
    for step in [
        "打开 设置 → 告警入库故障排查 → 重新检测；",
        "查看 Indexer 连接是否正常（HTTP 200）；",
        "对比「Indexer 最新告警」与「游标位置」：若已追平，表示暂无新告警，属正常；",
        "若游标为空且回溯窗口内无数据（如最新告警超过 24h），点击「扩大回溯至 7 天并重置游标」，再执行告警入库；",
        "确认 data/geoip 或 GEOLITE2_HOST_PATH 挂载正确（与 GeoIP 无关，但诊断页会一并显示）。",
    ]:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("9.3 常见手动运行返回码", level=2)
    add_table(
        doc,
        ["返回", "含义", "处理"],
        [
            ["skipped_lock", "Redis 锁被占用", "等待 2–5 分钟后重试"],
            ["skipped_budget", "今日调查预算用尽", "调大 investigation.daily_event_budget 或次日再跑"],
            ["skipped_not_configured", "协查 IM 未配置", "配置 WECOM_* 或跳过协查阶段"],
            ["errors 非空", "阶段执行异常", "查看 app 日志 docker compose logs app"],
        ],
    )

    doc.add_heading("9.4 Agent 调查失败", level=2)
    for t in [
        "检查 DEEPSEEK_API_KEY 是否有效、DEEPSEEK_BASE_URL 与 investigation.model_name 是否匹配；",
        "事件详情 → 调查轨迹查看具体 HTTP 错误；",
        "502/404 多为模型名或 API 地址配置错误。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("9.5 GeoIP 未就绪", level=2)
    doc.add_paragraph(
        "执行 docker compose exec app ma-geo-enrich --check 查看诊断。"
        "确保宿主机 ./data/geoip/GeoLite2-City.mmdb 为真实文件（非 Docker 误建的空目录），"
        "然后 docker compose up -d --force-recreate app。"
    )

    doc.add_heading("9.6 协查无响应", level=2)
    for t in [
        "确认 human_review.enabled=true 且 WECOM 凭据正确；",
        "企微回调 URL：https://<域名>/api/v1/human-review/webhooks/wecom 必须公网可达；",
        "也可在 Web 人工协查页直接回复。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("9.7 防御资产页为空", level=2)
    doc.add_paragraph(
        "白名单候选、画像建议来自已 concluded 且 Investigation 含 suggested_assets 的事件，"
        "需先完成 Agent 调查与防御资产沉淀阶段。"
        "entity_profiles 种子数据不会自动出现在「待确认白名单」列表。"
    )

    doc.add_heading("9.9 时区与入库异常", level=2)
    doc.add_paragraph(
        "Wazuh Indexer 告警 timestamp 可能带 +0800 时区，系统内部统一转为 UTC 存储。"
        "若出现 datetime offset-naive/aware 错误，需升级至已修复版本并重建 app 容器。"
    )

    doc.add_heading("9.10 nginx 构建失败（BOM）", level=2)
    doc.add_paragraph(
        "若 npm build 报 Unexpected token 或 package.json 解析错误，"
        "多为 UTF-8 BOM 导致。使用已包含 BOM 清理的 Dockerfile.nginx 重新 build nginx。"
    )

    doc.add_heading("9.11 防御资产与种子数据", level=2)
    doc.add_paragraph(
        "ma-entity-seed 写入 entity_profiles 与 on_duty_knowledge（如扫描器 IP），"
        "供 query_asset Skill 查询，不会直接出现在防御资产「待确认」列表。"
        "待确认项来自 Agent suggested_assets 经 defense_assets 流水线生成。"
    )

    doc.add_heading("9.8 日志与备份", level=2)
    for t in [
        "应用日志：docker compose logs -f app",
        "数据库备份：./deploy/scripts/backup-postgres.sh",
        "恢复：./deploy/scripts/restore-postgres.sh <备份文件>",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    # ===== 10 附录 =====
    doc.add_heading("10. 附录", level=1)
    doc.add_heading("10.1 API 手动运行端点（需管理员）", level=2)
    doc.add_paragraph("POST /api/v1/console/run/{pipeline}?limit=10")
    doc.add_paragraph(
        "pipeline 取值：ingestion、triage、aggregation、investigation、human_review、"
        "defense_assets、regression、health_check"
    )

    doc.add_heading("10.2 相关文档", level=2)
    for doc_name in [
        "docs/ingestion-phase2.md — 入库阶段设计",
        "docs/triage-phase3.md — 初筛阶段设计",
        "docs/ecs-deployment-phase10.md — ECS 部署 runbook",
        "docs/database-er.md — 数据库 ER 说明",
    ]:
        doc.add_paragraph(doc_name, style="List Bullet")

    doc.add_heading("10.3 技术支持联系", level=2)
    doc.add_paragraph("请保留本文档版本号与部署环境信息（Compose 文件、.env 非敏感项、诊断页截图）以便排查。")

    return doc


def main() -> None:
    base = Path(__file__).resolve().parent
    doc = build_document()
    paths = [
        base / "MA-MinSight-Operations-Manual.docx",
        base / "MA-MinSight-操作说明.docx",
    ]
    for out in paths:
        doc.save(out)
        print(f"Generated: {out}")


if __name__ == "__main__":
    main()
