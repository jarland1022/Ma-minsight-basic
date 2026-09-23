#!/usr/bin/env python3
"""Generate docs/MA-MinSight-开发者上手指南.docx from the developer onboarding guide."""

from __future__ import annotations

import shutil
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "MA-MinSight-开发者上手指南.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_code_block(doc: Document, text: str) -> None:
    for line in text.strip().splitlines():
        p = doc.add_paragraph(line)
        p.style = "No Spacing"
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val


def build_document() -> Document:
    doc = Document()
    title = doc.add_heading("MA-MinSight 开发者上手指南", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sub = doc.add_paragraph("代码结构解析 · 本地开发 · 修 Bug 方法论")
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    meta = doc.add_paragraph(f"文档日期：{date.today().isoformat()}")
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    add_para(
        doc,
        "读者：懂 Python，希望参与修 Bug / 小功能，不必先精通安全领域。",
        bold=True,
    )
    add_para(
        doc,
        "配套文档：运维排障见 docs/MA-MinSight-内部运维-runbook.md；表结构见 docs/database-er.md。",
    )

    # §1
    add_heading(doc, "1. 系统是做什么的？", 1)
    add_para(doc, "MA-MinSight 把 Wazuh 等来源的安全告警自动处理成可操作的调查结论：")
    add_code_block(
        doc,
        """告警入库 → 初筛分诊 → 事件聚合 → Agent(ReAct) 调查 → 人工协查(可选) → 防御资产沉淀
                ↓
            ARCHIVED（去重/低分/白名单，不进 Agent）""",
    )
    add_bullets(
        doc,
        [
            "初筛（Triage）：规则打分 + 路由，不是最终安全结论。",
            "Agent 调查：LLM + 多个 Skill（工具）查资产/情报/历史，输出结构化 verdict。",
            "Web 控制台：React 前端 + FastAPI 后端，给运维/业务人员看仪表盘、事件、协查 Inbox。",
        ],
    )
    add_para(doc, "单租户：一套 ECS 一套库，无 tenant_id。")

    # §2
    add_heading(doc, "2. 技术栈", 1)
    add_table(
        doc,
        ["层", "技术"],
        [
            ["后端", "Python 3.11+、FastAPI、SQLAlchemy 2 async、Alembic"],
            ["数据库", "PostgreSQL（+ pgvector 可选）"],
            ["缓存/锁", "Redis"],
            ["调度", "APScheduler（进程内，随 uvicorn 启动）"],
            ["LLM", "OpenAI 兼容 HTTP（DeepSeek / 百炼等）"],
            ["前端", "React 18、TypeScript、Vite、Ant Design、TanStack Query"],
            ["部署", "Docker Compose、Nginx 反代"],
            ["测试", "pytest、pytest-asyncio"],
        ],
    )

    # §3
    add_heading(doc, "3. 仓库目录结构", 1)
    add_code_block(
        doc,
        """AI-SEC/
├── app/                    # 后端 Python 包（核心）
│   ├── main.py             # FastAPI 入口 + 注册路由 + 启动调度
│   ├── core/               # 配置、Redis、日志、datetime 工具
│   ├── db/                 # Base、Session、枚举
│   ├── models/             # SQLAlchemy ORM（34 张表）
│   ├── api/routes/         # HTTP API（按域拆分）
│   ├── console/            # 仪表盘/事件列表等查询聚合
│   ├── services/           # 认证、system_config 等横切服务
│   ├── ingestion/          # 告警接入（Wazuh、Mock、HealthProbe）
│   ├── triage/             # 初筛流水线（Handler 链）
│   ├── aggregation/        # 告警 → Event 聚合
│   ├── agent/              # ReAct 循环、编排、结论、监督
│   ├── skills/             # Agent 可调用的 Skill（Function Calling）
│   ├── llm/                # LLM HTTP 客户端
│   ├── human_review/       # 人工协查、企微适配、过期协查单清理
│   ├── defense_assets/     # 判例、白名单候选、处置建议沉淀
│   ├── threat_intel/       # 外部情报同步 → threat_intel_entries
│   ├── eval/               # 回归测试、链路探针
│   ├── scheduler/          # 各阶段定时任务注册
│   ├── entity/             # 实体图（调查上下文）
│   ├── entity_seed/        # 资产画像种子数据 CLI
│   ├── geoip/              # GeoIP enrichment
│   └── tools/gateway/      # 外部 CMDB 等工具网关（可选）
├── web/                    # 前端 SPA
│   └── src/pages/          # Dashboard、Events、HumanReview 等
├── alembic/versions/       # 数据库迁移 001 … 017
├── tests/                  # pytest（按域分子目录）
├── deploy/                 # Docker、Nginx、脚本
├── docs/                   # 设计/阶段文档 + 本指南
├── docker-compose*.yml
├── pyproject.toml
└── README.md""",
    )

    # §4
    add_heading(doc, "4. 核心数据流（修 Bug 前先定位在哪一段）", 1)
    add_table(
        doc,
        ["阶段", "包路径", "调度", "手动触发 API"],
        [
            ["入库", "app/ingestion/", "scheduler/ingestion_jobs.py", "POST /api/v1/ingestion/..."],
            ["初筛", "app/triage/", "scheduler/triage_jobs.py", "POST /api/v1/console/run/triage"],
            ["聚合", "app/aggregation/", "scheduler/aggregation_jobs.py", ".../run/aggregation"],
            ["Agent", "app/agent/", "scheduler/investigation_jobs.py", ".../run/investigation"],
            ["协查", "app/human_review/", "scheduler/human_review_jobs.py", ".../run/human_review"],
            ["沉淀", "app/defense_assets/", "scheduler/defense_asset_jobs.py", ".../run/defense_assets"],
        ],
    )
    add_para(doc, "状态枚举（必读 app/db/enums.py）：", bold=True)
    add_bullets(
        doc,
        [
            "告警：NEW → TRIAGED / ARCHIVED → EVENT_LINKED",
            "事件：pending_review → investigating → human_pending / concluded",
            "PostgreSQL 存大写枚举名（如 ARCHIVED），Python 用 str, Enum 值为小写字符串，ORM 负责映射。",
        ],
    )

    # §5
    add_heading(doc, "5. 关键模块详解", 1)

    add_heading(doc, "5.1 配置：两层来源", 2)
    add_table(
        doc,
        ["来源", "文件", "用途"],
        [
            ["环境变量", "app/core/config.py（.env）", "DB、Redis、LLM Key、DEEPSEEK_*"],
            ["数据库", "system_config 表", "阈值、Agent 模型名、协查开关等"],
        ],
    )
    add_para(doc, "读取 DB 配置示例：")
    add_code_block(
        doc,
        """from app.services.system_config import get_config_value, get_config_int
from app.agent.config import load_investigation_config""",
    )
    add_para(
        doc,
        "注意：ensure_config_key() 只应在键缺失时插入默认值，不要覆盖运维在 Web 里改过的值"
        "（曾导致 investigation.model_name 被改回 deepseek-chat）。",
        bold=True,
    )
    add_para(doc, "Web 改配置：PUT /api/v1/system/config/{key}，前端 SettingsPage.tsx。")

    add_heading(doc, "5.2 初筛流水线（Handler 链）", 2)
    add_para(doc, "入口：app/triage/pipeline.py")
    add_code_block(
        doc,
        """DedupHandler → WhitelistHandler → ProfileHintHandler → RuleScoreHandler
  → LlmAssistHandler → RouteHandler → PersistHandler""",
    )
    add_bullets(
        doc,
        [
            "Dedup：60 分钟内同指纹 → ARCHIVE_DEDUP（仅重复条；首条继续打分）。",
            "Whitelist：唯一允许「规则直接归档」的路径（archive_whitelist）。",
            "RuleScore：severity*0.4 + frequency*0.3 + asset*0.2 + category*0.1。",
            "Router：默认 <30 低分归档，30–70 uncertain，≥70 deep_review。",
            "Persist：写 triage_results，更新 alerts.status。",
        ],
    )
    add_para(doc, "每个 Handler 实现 app/triage/handlers/base.py 的 handle(ctx, session, config)。")

    add_heading(doc, "5.3 Agent 调查（ReAct）", 2)
    add_code_block(
        doc,
        """InvestigationOrchestrator (orchestrator.py)
  → 选 pending_review 的 Event
  → build_investigation_context (context_builder.py)
  → AgentLoop.run (loop.py)
       → LLMClient.chat + tools
       → SkillRegistry.execute
       → submit_conclusion 或 degraded_conclusion
  → 写 investigation_conclusions
  → cancel_obsolete_sent_reviews (human_review/stale.py)""",
    )
    add_bullets(
        doc,
        [
            "模型名：来自 load_investigation_config().model_name（不是仅看 .env 的 DEEPSEEK_MODEL）。",
            "LLM 客户端：app/llm/client.py，POST {base_url}/chat/completions。",
            "Skill 注册：app/skills/bootstrap.py + app/skills/registry.py。",
            "失败降级：LLM 异常 → degraded_conclusion → 协查里出现「自动调查未完成：LLM error…」。",
        ],
    )

    add_heading(doc, "5.4 Skill（Agent 工具）", 2)
    add_para(doc, "基类：app/skills/base.py（Skill + SkillResult）")
    add_table(
        doc,
        ["Skill", "文件", "作用"],
        [
            ["query_asset", "query_asset.py", "查实体画像"],
            ["query_history_alerts", "query_history_alerts.py", "历史告警"],
            ["query_threat_intel", "query_threat_intel.py", "查 threat_intel_entries 缓存"],
            ["query_entity_graph", "query_entity_graph.py", "实体关系图"],
            ["simulate_disposition", "simulate_disposition.py", "处置推演（可配置关闭）"],
            ["submit_conclusion", "submit_conclusion.py", "提交最终结论（内置）"],
        ],
    )
    add_para(doc, "新增 Skill：实现类 → SkillRegistry.register → 在 bootstrap.py import。")

    add_heading(doc, "5.5 人工协查", 2)
    add_bullets(
        doc,
        [
            "派发：app/human_review/service.py 的 HumanReviewService.run",
            "清理旧单：app/human_review/stale.py（新调查完成后 / 协查派发前）",
            "回复触发复跑：reinvestigation.py + handle_reply",
        ],
    )

    add_heading(doc, "5.6 API 与前端", 2)
    add_bullets(
        doc,
        [
            "路由注册：app/main.py",
            "控制台手动跑流水线：app/api/routes/console.py → POST /api/v1/console/run/{pipeline}",
            "事件列表查询：app/console/queries.py",
            "前端路由：web/src/App.tsx；API 基址：web/src/api/client.ts",
        ],
    )
    add_para(doc, "改仅后端逻辑：往往只动 app/。改界面文案/列：动 web/src/。改表结构：ORM app/models/ + 新 Alembic revision。")

    add_heading(doc, "5.7 调度与分布式锁", 2)
    add_bullets(
        doc,
        [
            "调度器：app/scheduler/*.py，在 lifespan 里 start_scheduler()",
            "锁：app/agent/lock.py（investigation）、app/triage/lock.py 等，Redis SET NX EX",
        ],
    )

    # §6
    add_heading(doc, "6. 本地开发环境", 1)
    add_heading(doc, "6.1 准备", 2)
    add_code_block(
        doc,
        """git clone <repo> && cd AI-SEC
cp .env.example .env
python -m venv .venv
# Windows: .venv\\Scripts\\activate
pip install -e ".[dev]"
docker compose up -d postgres redis
alembic upgrade head
pytest tests/ -q""",
    )
    add_heading(doc, "6.2 启动", 2)
    add_code_block(
        doc,
        """# 终端 1：API + 内置调度
uvicorn app.main:app --reload --host 0.0.0.0 --port 8000

# 终端 2：前端
cd web && npm install && npm run dev
# http://localhost:5173  默认 admin / changeme""",
    )
    add_heading(doc, "6.3 Docker 全栈（接近生产）", 2)
    add_code_block(
        doc,
        """cp deploy/env/production.env.example .env
docker compose up -d --build""",
    )

    # §7
    add_heading(doc, "7. 如何修 Bug：推荐方法论", 1)
    add_heading(doc, "7.1 先分类问题", 2)
    add_table(
        doc,
        ["用户现象", "优先查"],
        [
            ["协查/调查 LLM 404", ".env、investigation.model_name、app/llm/client.py、容器内探活"],
            ["ReAct 某 Skill 报错", "app/skills/<name>.py、对应 tests/skills/"],
            ["ARCHIVED 数量异常", "app/triage/handlers/dedup.py、router.py、triage_results 表"],
            ["Inbox 文案不更新", "是否旧 human_review_requests；human_review/stale.py"],
            ["调度锁占用", "Redis investigate:lock:global、agent/lock.py"],
            ["前端 build 失败", "web/src/、npm run build"],
            ["迁移失败", "alembic/versions/，revision id ≤32 字符"],
        ],
    )

    add_heading(doc, "7.2 调试路径", 2)
    add_numbered(
        doc,
        [
            "Web 设置 → 手动运行对应流水线，看返回 JSON（processed、errors、skipped_lock）。",
            "查库：优先 docker compose exec app python + SQLAlchemy，避免 psql 枚举/用户搞错。",
            "看日志：docker compose logs app --tail 200。",
            "写/跑测试：在 tests/ 加最小复现，pytest tests/path/test_x.py -q。",
            "小步提交：一个 Bug 一个 PR，附带「现象 → 根因 → 验证步骤」。",
        ],
    )

    add_heading(doc, "7.3 常见坑（本项目真实踩过）", 2)
    add_table(
        doc,
        ["坑", "正确做法"],
        [
            ["改 .env 后无效", "docker compose up -d --force-recreate app，不是 restart"],
            ["Agent 模型不对", "改 system_config + 确认 ensure_config_key 不覆盖"],
            ["PostgreSQL 时间比较", "用 utc_now() / as_naive_utc()，禁止 datetime.now(UTC) 直接绑 TIMESTAMP WITHOUT TIME ZONE"],
            ["SQL 枚举", "库内为大写 ARCHIVED；用 status::text 或 ORM"],
            ["只重建 nginx", "后端 Python 改动必须重建 app 镜像"],
            ["协查 question 是快照", "改配置不会改旧单；需重跑调查 + 协查派发"],
        ],
    )

    add_heading(doc, "7.4 修 Bug 最小流程示例", 2)
    add_para(doc, "例：Skill 查库时区错误", bold=True)
    add_numbered(
        doc,
        [
            "读 Trace 里 skill 名 → 打开 app/skills/query_threat_intel.py",
            "确认 expires_at >= now 的 now 是否 naive UTC",
            "改 as_naive_utc(utc_now())，补 tests/skills/test_query_threat_intel_timezone.py",
            "pytest tests/skills/ -q",
            "部署：docker compose ... up -d --build --force-recreate app",
        ],
    )

    # §8
    add_heading(doc, "8. 测试", 1)
    add_code_block(
        doc,
        """pytest tests/ -q              # 全量
pytest tests/triage/ -q         # 初筛
pytest tests/agent/ -q          # Agent
pytest tests/skills/ -q         # Skill
pytest tests/eval/ -q           # 回归/探针""",
    )
    add_bullets(
        doc,
        [
            "异步测试：pytest.ini 里 asyncio_mode = auto",
            "回归夹具：app/eval/regression/fixtures/*.json",
            "探针：app/eval/health_check/",
            "修 Agent/初筛逻辑后：Web 评测 → 回归测试 / 链路探针 跑一轮",
        ],
    )

    # §9
    add_heading(doc, "9. 数据库迁移", 1)
    add_code_block(
        doc,
        """alembic revision -m "short_desc"   # revision id 保持 ≤32 字符！
alembic upgrade head
docker compose exec app alembic upgrade head""",
    )
    add_para(doc, "模型定义：app/models/*.py，汇总在 app/models/__init__.py。")

    # §10
    add_heading(doc, "10. 前端入门（会 Python 的开发者）", 1)
    add_table(
        doc,
        ["任务", "位置"],
        [
            ["加列表列", "web/src/pages/EventsPage.tsx + web/src/types/api.ts"],
            ["改 API 字段", "后端 app/console/queries.py 或 routes，再改 TS 类型"],
            ["用户指引", "web/src/content/userGuide.ts"],
            ["构建验证", "cd web && npm run build（Docker 镜像里也会跑）"],
        ],
    )
    add_para(
        doc,
        "Ant Design Table 排序列需注意 TypeScript：SortOrder 类型用 'ascend' | 'descend'，见 EventsPage.tsx。",
    )

    # §11
    add_heading(doc, "11. 建议阅读顺序（3 天入门）", 1)
    add_heading(doc, "第 1 天：跑通 + 数据流", 2)
    add_numbered(
        doc,
        [
            "README.md",
            "本指南 §3–§4",
            "docs/database-er.md（扫表关系）",
            "本地或 Docker 跑起来，Web 点一遍手动运行流水线",
        ],
    )
    add_heading(doc, "第 2 天：初筛 + Agent", 2)
    add_numbered(
        doc,
        [
            "app/triage/pipeline.py + 各 handlers/",
            "app/agent/orchestrator.py → loop.py",
            "app/skills/query_asset.py（Skill 范例）",
            "app/llm/client.py",
        ],
    )
    add_heading(doc, "第 3 天：协查 + 修 Bug 练习", 2)
    add_numbered(
        doc,
        [
            "app/human_review/service.py + stale.py",
            "app/services/system_config.py",
            "tests/triage/test_dedup.py（测试范例）",
            "读 MA-MinSight-内部运维-runbook.md §7 ARCHIVED 审计",
        ],
    )
    add_para(doc, "按需深读阶段文档：docs/triage-phase3.md、investigation-phase5.md、human-review-phase6.md 等。")

    # §12
    add_heading(doc, "12. 代码规范与协作", 1)
    add_bullets(
        doc,
        [
            "Python：类型标注、async/await 与现有文件保持一致；少写大而全的抽象。",
            "配置默认值：用 ensure_config_key 仅插入缺失键，禁止静默覆盖生产配置。",
            "时间：凡写 PostgreSQL TIMESTAMP WITHOUT TIME ZONE，统一 utc_now() / as_naive_utc()。",
            "提交：说明「为什么」；.env、密钥、证书勿提交。",
            "PR 自检：pytest 相关目录 + 若动前端则 npm run build。",
        ],
    )

    # §13
    add_heading(doc, "13. 快速定位表（grep 关键词）", 1)
    add_table(
        doc,
        ["关键词", "可能文件"],
        [
            ["route_decision", "app/triage/handlers/router.py"],
            ["degraded_conclusion", "app/agent/conclusion.py, loop.py"],
            ["investigation.model_name", "app/agent/config.py"],
            ["HumanReviewRequest", "app/human_review/service.py, models/human_review.py"],
            ["ARCHIVE_DEDUP", "app/triage/handlers/dedup.py"],
            ["list_events", "app/console/queries.py"],
            ["manual_run", "app/api/routes/console.py"],
            ["SkillRegistry", "app/skills/registry.py"],
        ],
    )

    # §14
    add_heading(doc, "14. 第一条练手任务（可选）", 1)
    add_numbered(
        doc,
        [
            "在本地跑通 §6。",
            "读 tests/triage/test_dedup.py，理解 dedup 行为。",
            "用 app 容器 SQL 验证：ARCHIVED 是否几乎全是 ARCHIVE_DEDUP（见内部 runbook §7）。",
            "故意在 query_threat_intel.py 写错时区，看 Trace 报错，再改回 as_naive_utc，跑 pytest tests/skills/ -q。",
        ],
    )
    add_para(
        doc,
        "完成以上即具备独立修 Skill/初筛/配置类 Bug 的基础；Agent 提示词、评测、前端需再结合对应 docs/*-phase*.md 深入。",
    )

    # §15
    add_heading(doc, "15. 相关文档索引", 1)
    add_table(
        doc,
        ["文档", "内容"],
        [
            ["MA-MinSight-内部运维-runbook.md", "ECS 部署、LLM、协查、SQL 审计"],
            ["database-er.md", "34 表 ER、约束"],
            ["ecs-deployment-phase10.md", "生产部署"],
            ["triage-phase3.md", "初筛设计"],
            ["investigation-phase5.md", "Agent 设计"],
            ["human-review-phase6.md", "协查设计"],
            ["regression-health-phase9.md", "回归与探针"],
        ],
    )

    add_heading(doc, "附录：一句话定位", 1)
    add_para(
        doc,
        "MA-MinSight 是 FastAPI + Handler 链初筛 + ReAct Agent(Skill)；"
        "修 Bug 先确定在流水线哪一段，再查 app/db/enums.py 和 system_config，"
        "时间字段统一用 utc_now()，改 Python 后重建 app 容器。",
        bold=True,
    )

    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    return doc


def main() -> None:
    doc = build_document()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = Path(tmp.name)
    try:
        doc.save(tmp_path)
        shutil.copy2(tmp_path, OUTPUT)
        print(f"Wrote {OUTPUT}")
    except PermissionError:
        fallback = OUTPUT.with_name("MA-MinSight-开发者上手指南-v1.docx")
        shutil.copy2(tmp_path, fallback)
        print(f"Target locked; wrote {fallback} (close Word and rename to replace original)")
    finally:
        tmp_path.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
